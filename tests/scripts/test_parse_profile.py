"""Small, self-contained workbooks exercise regeneration without SDK downloads."""

import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

from scripts.parse_profile import MESSAGE_COLUMNS, TYPE_COLUMNS, render


@pytest.fixture
def workbook(tmp_path):
    book = Workbook()
    types = book.active
    types.title = "Types"
    types.append(TYPE_COLUMNS)
    for row in [
        ("mesg_num", "uint16"),
        ("", "", "alpha", "1"),
        ("", "", "beta", "2"),
        ("", "", "pad", "105"),
        ("selector", "enum"),
        ("", "", "deprecated", "0x01"),
        ("", "", "enabled", "1"),
    ]:
        types.append(row)
    messages = book.create_sheet("Messages")
    messages.append(MESSAGE_COLUMNS)

    def row(**values):
        messages.append([values.get(column, "") for column in MESSAGE_COLUMNS])

    row(**{"Field Type": "EXAMPLE MESSAGES"})
    for name, ref_num in (("alpha", 7), ("beta", 9)):
        row(**{"Message Name": name})
        row(**{"Field Def #": "0", "Field Name": "packed", "Field Type": "uint16"})
        row(
            **{
                "Field Name": "variant",
                "Field Type": "uint16",
                "Components": "total",
                "Bits": "8",
                "Scale": "1.024",
                "Accumulate": "1",
                "Ref Field Name": "kind,kind",
                "Ref Field Value": "deprecated,enabled",
            }
        )
        row(
            **{
                "Field Def #": "1",
                "Field Name": "total",
                "Field Type": "uint32",
                "Scale": "1.024",
            }
        )
        # The selector deliberately follows its references and differs per message.
        row(
            **{
                "Field Def #": str(ref_num),
                "Field Name": "kind",
                "Field Type": "selector",
            }
        )
        row(
            **{
                "Field Def #": "2",
                "Field Name": "samples",
                "Field Type": "uint8",
                "Array": "[N]",
            }
        )
        row(
            **{
                "Field Def #": "3",
                "Field Name": "fixed",
                "Field Type": "uint8",
                "Array": "[3]",
            }
        )
    path = tmp_path / "Profile.xlsx"
    book.save(path)
    return book, path


def namespace(outputs, filename):
    from fittie.profile.field_type import FieldType, FieldTypeValue
    from fittie.profile.message_profile import FieldProfile, MessageProfile, SubField

    result = dict(
        FieldType=FieldType,
        FieldTypeValue=FieldTypeValue,
        FieldProfile=FieldProfile,
        MessageProfile=MessageProfile,
        SubField=SubField,
    )
    source = "\n".join(
        line for line in outputs[filename].splitlines() if not line.startswith("from .")
    )
    exec(source, result)
    return result


def test_self_contained_generation_and_message_local_references(workbook):
    _, path = workbook
    output = render(path, "21.217.0")
    assert output == render(path, "21.217.0")
    messages = namespace(output, "messages.py")["MESSAGES"]
    for num, ref_num in ((1, 7), (2, 9)):
        fields = messages[num].fields
        assert not fields[0].is_array
        assert fields[2].is_array and fields[2].array == "N"
        assert fields[3].array == 3
        sub = fields[0].subfields[0]
        assert sub.accumulate == 1 and sub.scale == 1.024
        assert [ref["field_number"] for ref in sub.refs] == [ref_num, ref_num]
        assert [ref["value_number"] for ref in sub.refs] == [1, 1]
    assert messages[105].fields == {}
    assert (
        namespace(output, "fit_types.py")["FIT_TYPES"]["selector"].values[1].value_name
        == "enabled"
    )
    assert json.loads(output["source.json"])["messages"] == 3


@pytest.mark.parametrize(
    "column,value,error",
    [
        ("Ref Field Name", "absent,kind", "unknown reference field"),
        ("Ref Field Value", "missing,enabled", "unresolved reference"),
        ("Ref Field Value", "enabled", "reference names and values"),
        ("Components", "absent", "unknown component target"),
        ("Components", "total,total", ""),
        ("Bits", "8,8,8", "inconsistent bits"),
        ("Bits", "65", "bit width"),
        ("Scale", "0", "nonzero"),
        ("Scale", "nan", "non-finite"),
        ("Array", "[0]", "invalid array"),
        ("Accumulate", "2", "accumulate"),
        ("Field Type", "unknown", "unknown field type"),
        ("Field Name", "=1+1", "formulas"),
    ],
)
def test_bad_metadata_is_rejected(workbook, column, value, error):
    book, path = workbook
    book["Messages"].cell(5, MESSAGE_COLUMNS.index(column) + 1, value)
    book.save(path)
    if error:
        with pytest.raises(ValueError, match=error):
            render(path, "21.217.0")
    else:
        # A single bit width is valid for repeated components.
        assert render(path, "21.217.0")


def test_cli_bootstraps_without_package_and_checks_without_writing(workbook, tmp_path):
    _, path = workbook
    script = Path(__file__).resolve().parents[2] / "scripts" / "parse_profile.py"
    isolated = tmp_path / "parse_profile.py"
    isolated.write_bytes(script.read_bytes())
    args = [
        sys.executable,
        "-I",
        str(isolated),
        str(path),
        "--version",
        "21.217.0",
        "--output-root",
        str(tmp_path),
    ]
    subprocess.run(args, cwd=tmp_path, check=True, capture_output=True)
    subprocess.run([*args, "--check"], cwd=tmp_path, check=True, capture_output=True)
    output = tmp_path / "fittie" / "profile" / "mesg_nums.py"
    output.write_text("stale\n")
    assert subprocess.run([*args, "--check"], capture_output=True).returncode == 1
    assert output.read_text() == "stale\n"


def test_invalid_workbook_leaves_existing_outputs_untouched(workbook, tmp_path):
    book, path = workbook
    book["Messages"].cell(5, MESSAGE_COLUMNS.index("Scale") + 1, "0")
    book.save(path)
    script = Path(__file__).resolve().parents[2] / "scripts" / "parse_profile.py"
    output = tmp_path / "fittie" / "profile"
    output.mkdir(parents=True)
    sentinel = output / "fit_types.py"
    sentinel.write_text("keep me\n")
    result = subprocess.run(
        [
            sys.executable,
            str(script),
            str(path),
            "--version",
            "21.217.0",
            "--output-root",
            str(tmp_path),
        ],
        capture_output=True,
    )
    assert result.returncode == 2
    assert sentinel.read_text() == "keep me\n"
    assert list(output.iterdir()) == [sentinel]
